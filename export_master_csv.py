import openpyxl
import csv
from datetime import datetime

wb = openpyxl.load_workbook('Biocare Dispatch Tracker.xlsx', data_only=True)
all_records = []

for name in wb.sheetnames:
    if '-' in name and len(name) == 10:
        ws = wb[name]
        for row in range(4, 104):
            order = str(ws.cell(row=row, column=2).value or '').strip()
            if order and order != 'None':
                raw_date = ws.cell(row=row, column=1).value or name
                if isinstance(raw_date, datetime):
                    date_str = raw_date.strftime('%d-%m-%Y')
                else:
                    date_str = str(raw_date).strip().split(' ')[0]
                    # If format is YYYY-MM-DD convert to DD-MM-YYYY
                    if len(date_str) == 10 and date_str[4] == '-' and date_str[7] == '-':
                        parts = date_str.split('-')
                        date_str = f"{parts[2]}-{parts[1]}-{parts[0]}"
                
                cust = str(ws.cell(row=row, column=3).value or '').strip()
                status = str(ws.cell(row=row, column=4).value or 'Dispatched').strip()
                remarks = str(ws.cell(row=row, column=5).value or '').strip()
                if remarks == 'None': remarks = ''

                all_records.append({
                    'dispatch_date': date_str,
                    'order_no': order,
                    'customer': cust,
                    'status': status,
                    'remarks': remarks
                })

# Save to CSV
csv_filename = 'biocare_dispatches_master.csv'
with open(csv_filename, 'w', newline='', encoding='utf-8') as f:
    writer = csv.DictWriter(f, fieldnames=['dispatch_date', 'order_no', 'customer', 'status', 'remarks'])
    writer.writeheader()
    writer.writerows(all_records)

print(f"Exported {len(all_records)} records to {csv_filename} successfully!")
